"""
QRadar Rule Status Report — v3

MAJOR CHANGE FROM v2: offense data is no longer used anywhere in this script. Two reasons drove
this, both from the same root cause (short offense retention):
  1. "First/Last Triggered" is now pulled directly from QRadar Use Case Manager (UCM)'s Use Case
     Explorer report API, which tracks these independently of offense retention.
  2. The "Event is part of an offense" rule filter isn't exposed via the core /api/analytics/rules
     endpoint at all — it lives in UCM's own data model. UCM's Use Case Explorer supports it
     natively as a documented filter, which is also why the old text-search-based filter attempt
     never found anything: it was searching a payload that structurally couldn't contain it.

Because offense retention was the whole reason to stop trusting offense-derived "Dead" too, "Dead"
is now ALSO redefined as "never triggered, per UCM" rather than "zero offenses in retained
history" — this keeps the whole report internally consistent on one underlying data source. Flag
this if you wanted Dead to stay offense-based for some other reason.

IMPORTANT — please read before running:
  UCM's REST API is reached through the QRadar app-proxy path, which needs your Use Case Manager
  app's numeric ID (UCM_APP_ID) and, per IBM's docs, an Authorized Service token sent via a SEC
  header (UCM_AUTH_TOKEN) rather than the Basic Auth used elsewhere in this script. I have solid
  documentation for the REQUEST shapes (filters, column codes, the generate -> poll -> fetch
  workflow) but not a verified example of the exact RESPONSE shape for the status/result calls in
  your specific app version — those can vary. Set UCM_DEBUG = True for your first run: it prints
  the raw JSON from every UCM call so you (or I, if you paste it back) can quickly confirm or
  correct the few assumptions marked below with "ASSUMPTION:".

  The most bulletproof way to get filters/columns that are guaranteed correct for YOUR app version:
  open Use Case Explorer in the QRadar UI, set the filters by hand (Rule enabled = Yes, Action =
  "Ensure the detected event is part of an offense"), open your browser's dev tools Network tab,
  and copy the exact POST body QRadar sends to /api/use_case_explorer. Paste that over the
  `filters` list in fetch_ucm_rule_report() below if it differs from what's here.
"""

import os
import json
import shutil
import time
import tempfile
import traceback
import logging
from datetime import datetime, timedelta

import requests
import urllib3
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import win32com.client

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

logging.basicConfig(
    level=logging.WARNING,
    format='%(asctime)s [%(levelname)s] %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
logger = logging.getLogger(__name__)


QRADAR_HOST     = 'https://your-qradar-host'
QRADAR_USERNAME = 'your-username'   # used for the /api/help/versions connection check, and as a
QRADAR_PASSWORD = 'your-password'   # fallback UCM auth if UCM_AUTH_TOKEN is left blank
VERIFY_SSL      = False

# --- QRadar Use Case Manager (UCM) app ---------------------------------------
UCM_APP_ID     = ''   # from the browser URL when you open Use Case Manager in the QRadar UI,
                       # e.g. https://<host>/console/plugins/<APP_ID>/app_proxy/...
UCM_AUTH_TOKEN = ''    # An Authorized Service Token (QRadar Admin > Authorized Services), sent via
                       # the SEC header — IBM's docs specify this for UCM's API. Leave blank to
                       # fall back to Basic Auth with QRADAR_USERNAME/PASSWORD instead.
UCM_POLL_INTERVAL_SECS = 2
UCM_POLL_TIMEOUT_SECS  = 120
UCM_DEBUG = False      # True: print raw JSON from every UCM API call. Turn on for your first run.

# Report columns requested from Use Case Explorer (see IBM's "Report column codes" reference).
# ASSUMPTION: result rows come back keyed by these same short codes. _normalize_ucm_row() also
# handles the case where a row is a plain list in this same column order, and _row_get() also
# tries a couple of plausible full-name alternates — but if UCM_DEBUG shows a genuinely different
# shape, this is the one place to fix it.
UCM_REPORT_COLUMNS = ["ID", "N", "EN", "T", "RO", "GR", "A", "RE", "AD", "RED", "CD", "MD",
                      "FTG", "LTG", "TG", "NO"]

# --- Time windows -------------------------------------------------------------
NEWLY_TRIGGERED_WINDOW_DAYS    = 30   # "first triggered within N days" -> newly validated rule
TRACKER_VALIDATION_WINDOW_DAYS = 7    # "last triggered within N days" -> tracker auto-validation

# --- Tracker file --------------------------------------------------------------
TRACKER_EXCEL_PATH            = ''
TRACKER_SHEET_NAME            = None      # None = first sheet in the workbook; or an exact name
TRACKER_HEADER_ROW            = 1
AUTO_UPDATE_TRACKER_TRIGGERED = True
TRACKER_BACKUP_BEFORE_WRITE   = True

OUTPUT_DIR = r'C:\path\to\your\output'

REPORT_TITLE = 'Rules Under Investigation & Dead'

EMAIL_TABLE_ROW_CAP = 25
EMAIL_HIGH_IMPORTANCE_DEAD_THRESHOLD = 15
DEAD_RULE_STALE_AGE_DAYS = 365

REQUEST_TIMEOUT    = 30
MAX_RETRIES        = 3
RETRY_DELAY_BASE   = 1.5
API_PAGE_SIZE      = 9999
MAX_PAGES          = 150
QRADAR_API_VERSION = '14.0'

OUTPUT_EXCEL = os.path.join(OUTPUT_DIR, 'qradar_rules_under_investigation_and_dead.xlsx')

_MAPI_PR_ATTACH_CONTENT_ID = "http://schemas.microsoft.com/mapi/proptag/0x3712001F"

_RULE_NAME_COL_ALIASES = ['alert rule name', 'rule name', 'name']
_TRIGGERED_COL_ALIASES = ['triggered', 'fired', 'confirmed', 'validated']
_TRUE_STRINGS = {'true', 'yes', 'y', '1', 'x', '✓', 'done', 'complete', 'completed'}


def _api_get_page(path, range_start, range_end, params=None, label='request'):
    url     = f"{QRADAR_HOST.rstrip('/')}{path}"
    headers = {
        'Accept':  'application/json',
        'Version': QRADAR_API_VERSION,
        'Range':   f'items={range_start}-{range_end}',
    }
    last_err = None

    for attempt in range(MAX_RETRIES):
        if attempt > 0:
            wait = RETRY_DELAY_BASE * (2 ** (attempt - 1))
            logger.warning("Retry %d/%d for '%s' — waiting %.1fs after %s",
                           attempt, MAX_RETRIES - 1, label, wait,
                           type(last_err).__name__)
            time.sleep(wait)
        try:
            resp = requests.get(
                url,
                params=params,
                auth=(QRADAR_USERNAME, QRADAR_PASSWORD),
                verify=VERIFY_SSL,
                timeout=REQUEST_TIMEOUT,
                headers=headers,
            )
            if resp.status_code in (200, 206):
                total = None
                cr = resp.headers.get('Content-Range', '')
                if cr:
                    try:
                        total = int(cr.split('/')[-1].strip())
                    except Exception:
                        pass
                return resp.json(), total
            elif resp.status_code == 401:
                raise RuntimeError("Authentication failed — check QRADAR_USERNAME / PASSWORD")
            else:
                last_err = RuntimeError(f"HTTP {resp.status_code}")
                logger.warning("HTTP %d for '%s' (attempt %d/%d)",
                               resp.status_code, label, attempt + 1, MAX_RETRIES)
        except (requests.exceptions.Timeout, requests.exceptions.ConnectionError) as exc:
            last_err = exc
            logger.warning("%s on attempt %d for '%s'",
                           type(exc).__name__, attempt + 1, label)
        except RuntimeError:
            raise
        except Exception:
            logger.error("Non-retriable error for '%s':\n%s", label, traceback.format_exc())
            raise

    raise RuntimeError(f"All {MAX_RETRIES} attempts failed for '{label}': {last_err}")


def _api_get(path, params=None, label='request'):
    items, total = _api_get_page(path, 0, API_PAGE_SIZE, params=params, label=label)
    if total is not None and total > API_PAGE_SIZE + 1:
        logger.warning(
            "Pagination cap hit for '%s': %d items on server, only %d fetched.",
            label, total, API_PAGE_SIZE + 1
        )
    return items


def validate_config():
    problems = []

    if QRADAR_HOST.rstrip('/') in ('https://your-qradar-host', ''):
        problems.append("QRADAR_HOST is still the placeholder — set it to your real console URL.")
    if not UCM_APP_ID:
        problems.append(
            "UCM_APP_ID is blank. Open Use Case Manager in the QRadar UI and copy the numeric "
            "app ID from the URL: https://<host>/console/plugins/<APP_ID>/app_proxy/..."
        )
    if not UCM_AUTH_TOKEN and (QRADAR_USERNAME in ('your-username', '') or QRADAR_PASSWORD in ('your-password', '')):
        problems.append(
            "UCM_AUTH_TOKEN is blank and QRADAR_USERNAME/PASSWORD look like placeholders — "
            "UCM calls will have no valid auth either way."
        )
    if OUTPUT_DIR == r'C:\path\to\your\output':
        problems.append("OUTPUT_DIR is still the placeholder path.")

    if TRACKER_VALIDATION_WINDOW_DAYS <= 0:
        problems.append("TRACKER_VALIDATION_WINDOW_DAYS must be a positive number of days.")
    if NEWLY_TRIGGERED_WINDOW_DAYS <= 0:
        problems.append("NEWLY_TRIGGERED_WINDOW_DAYS must be a positive number of days.")

    if TRACKER_EXCEL_PATH and not os.path.exists(TRACKER_EXCEL_PATH):
        problems.append(
            f"TRACKER_EXCEL_PATH is set to '{TRACKER_EXCEL_PATH}' but that file doesn't exist yet."
        )

    if problems:
        print("⚠  Configuration issues found:")
        for p in problems:
            print(f"   - {p}")
        print()

    fatal = any(
        ('placeholder' in p and 'QRADAR_' in p) or 'must be' in p or 'UCM_APP_ID' in p
        for p in problems
    )
    return not fatal


def test_connection():
    print("🔗 Testing base QRadar connection (Basic Auth)...")
    try:
        result = _api_get('/api/help/versions', label='connection test')
        if result:
            print("✅ QRadar connection successful!")
            return True
        print("⚠  Unexpected empty response from /api/help/versions")
        return False
    except RuntimeError as e:
        print(f"❌ {e}")
        return False
    except Exception as e:
        print(f"❌ Connection failed: {e}")
        return False


# --- Use Case Manager (UCM) -------------------------------------------------

def _ucm_base_url():
    return f"{QRADAR_HOST.rstrip('/')}/console/plugins/{UCM_APP_ID}/app_proxy/api"


def _ucm_request_kwargs():
    headers = {'Content-Type': 'application/json', 'Accept': 'application/json'}
    kwargs = {'verify': VERIFY_SSL, 'timeout': REQUEST_TIMEOUT, 'headers': headers}
    if UCM_AUTH_TOKEN:
        headers['SEC'] = UCM_AUTH_TOKEN
    else:
        kwargs['auth'] = (QRADAR_USERNAME, QRADAR_PASSWORD)
    return kwargs


def _ucm_debug_print(label, payload):
    if not UCM_DEBUG:
        return
    try:
        text = json.dumps(payload, default=str)
    except Exception:
        text = str(payload)
    if len(text) > 2000:
        text = text[:2000] + '…'
    print(f"\n🔎 UCM_DEBUG — {label}:\n{text}\n")


def _row_get(row, *keys):
    for k in keys:
        if k in row and row[k] not in (None, ''):
            return row[k]
    return None


def _normalize_ucm_row(row, columns):
    if isinstance(row, dict):
        return row
    if isinstance(row, (list, tuple)):
        return dict(zip(columns, row))
    return {}


def _parse_ucm_date(val):
    """
    Best-effort parse of a UCM First/Last/Creation/Modification date value. ASSUMPTION: could be
    epoch milliseconds (like the rest of QRadar's API) or a formatted date string; both are
    handled. A rule that has never triggered should come back empty/None here, not a real date.
    Returns (epoch_ms_or_None, display_string).
    """
    if val in (None, '', 'Never', 'N/A', '-', 0, '0'):
        return None, '—'
    if isinstance(val, (int, float)) or (isinstance(val, str) and val.strip().isdigit()):
        try:
            ms = int(val)
            return ms, datetime.fromtimestamp(ms / 1000).strftime('%Y-%m-%d')
        except (ValueError, OSError, OverflowError):
            pass
    for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%dT%H:%M:%S', '%m/%d/%Y %H:%M:%S',
                '%b %d, %Y %I:%M:%S %p', '%Y-%m-%d'):
        try:
            dt = datetime.strptime(str(val).strip(), fmt)
            return int(dt.timestamp() * 1000), dt.strftime('%Y-%m-%d')
        except ValueError:
            continue
    logger.warning(
        "Could not parse UCM date value %r — leaving it blank. Set UCM_DEBUG=True, check the raw "
        "value's format, and add it to _parse_ucm_date().", val
    )
    return None, '—'


def _parse_ucm_triggered_count(val):
    if val in (None, ''):
        return 0
    try:
        return int(val)
    except (TypeError, ValueError):
        return 1 if _to_bool(val) else 0


def fetch_ucm_rule_report():
    """
    3-step Use Case Explorer workflow: POST to generate -> GET status until complete -> GET result
    (paginated). Filters ask UCM to do the "enabled" + "not a building block" + "action = part of
    an offense" filtering server-side, so no separate /api/analytics/rules call or manual JSON
    text-search is needed anymore.

    ASSUMPTION: the "enabled" filter below follows the same documented pattern as IBM's "action"/
    "response"/"group" examples, but IBM's own filter reference page didn't give a literal example
    for "enabled" specifically — if this comes back with 0 or clearly-wrong rows, get the exact
    filter from your browser's dev tools as described in the module docstring above.
    """
    filters = [
        {"name": "enabled", "type": "ATTRIBUTE", "recursive": True,
         "matchCriteria": "PARTIAL", "values": [True], "attributeName": "", "valueType": "COMMON"},
        {"name": "rule", "type": "ATTRIBUTE", "recursive": True,
         "matchCriteria": "PARTIAL", "values": [True], "attributeName": "", "valueType": "EXCLUSIVE_COMMON"},
        {"name": "action", "type": "ATTRIBUTE", "recursive": True,
         "matchCriteria": "PARTIAL", "values": ["offense"], "attributeName": "", "valueType": "COMMON"},
    ]
    body = {"filters": filters, "columns": UCM_REPORT_COLUMNS}
    base = _ucm_base_url()

    print("📥 Requesting Use Case Explorer report (enabled + not-a-building-block + "
          "'part of an offense' action)...")
    post_resp = requests.post(f"{base}/use_case_explorer", json=body, **_ucm_request_kwargs())
    if post_resp.status_code not in (200, 201, 202):
        raise RuntimeError(f"UCM report generation failed: HTTP {post_resp.status_code} — "
                            f"{post_resp.text[:300]}")
    post_json = post_resp.json()
    _ucm_debug_print("POST /use_case_explorer response", post_json)

    report_id = post_json.get('reportId') or post_json.get('id') or post_json.get('report_id')
    if not report_id:
        raise RuntimeError(f"UCM report generation didn't return a reportId: {post_resp.text[:300]}")

    # Poll for completion. ASSUMPTION: a top-level "status" string that contains "COMPLETE" on
    # success or "FAIL"/"ERROR" on failure — confirm against UCM_DEBUG output.
    deadline = time.time() + UCM_POLL_TIMEOUT_SECS
    status = None
    while time.time() < deadline:
        st_resp = requests.get(f"{base}/use_case_explorer/{report_id}/status", **_ucm_request_kwargs())
        st_resp.raise_for_status()
        st_json = st_resp.json()
        _ucm_debug_print("GET .../status response", st_json)
        status = str(st_json.get('status', '')).upper()
        if 'COMPLETE' in status or 'SUCCESS' in status or 'DONE' in status:
            break
        if 'FAIL' in status or 'ERROR' in status:
            raise RuntimeError(f"UCM report generation failed with status: {status}")
        time.sleep(UCM_POLL_INTERVAL_SECS)
    else:
        raise RuntimeError(f"UCM report did not complete within {UCM_POLL_TIMEOUT_SECS}s "
                            f"(last status: {status})")

    # Paginate through results. ASSUMPTION: a `page` query param, 0-indexed, empty list = done.
    all_rows = []
    page = 0
    while True:
        res_resp = requests.get(f"{base}/use_case_explorer/{report_id}/result",
                                 params={'page': page}, **_ucm_request_kwargs())
        res_resp.raise_for_status()
        payload = res_resp.json()
        if page == 0:
            _ucm_debug_print("GET .../result response (page 0)", payload)

        if isinstance(payload, list):
            page_rows = payload
        elif isinstance(payload, dict):
            page_rows = payload.get('results') or payload.get('data') or payload.get('rows') or []
        else:
            page_rows = []

        if not page_rows:
            break
        all_rows.extend(page_rows)
        page += 1
        if page > 1000:
            logger.warning("UCM result pagination exceeded 1000 pages — stopping early; "
                            "data may be incomplete.")
            break

    print(f"   ✅ {len(all_rows)} rule(s) returned by Use Case Explorer.")
    return all_rows


def _html_escape(text):
    if text is None:
        return ''
    text = str(text)
    return (text.replace('&', '&amp;')
                .replace('<', '&lt;')
                .replace('>', '&gt;')
                .replace('"', '&quot;'))


def _csv_formula_safe(text):
    if text is None:
        return ''
    text = str(text)
    if text and text[0] in ('=', '+', '-', '@', '\t', '\r'):
        return "'" + text
    return text


def _native(val):
    try:
        if pd.isna(val):
            return ''
    except (TypeError, ValueError):
        pass
    if isinstance(val, pd.Timestamp):
        return val.to_pydatetime()
    if hasattr(val, 'item'):
        try:
            return val.item()
        except Exception:
            pass
    return val


def _age_days_from_epoch(val):
    if not val:
        return None
    try:
        created = datetime.fromtimestamp(int(val) / 1000)
        return (datetime.now() - created).days
    except (TypeError, ValueError, OSError, OverflowError):
        return None


def _find_col(columns, aliases):
    lower_map = {str(c).strip().lower(): c for c in columns}
    for alias in aliases:
        if alias in lower_map:
            return lower_map[alias]
    return None


def _to_bool(val):
    try:
        if pd.isna(val):
            return False
    except (TypeError, ValueError):
        pass
    if isinstance(val, bool):
        return val
    if isinstance(val, (int, float)):
        return val != 0
    return str(val).strip().lower() in _TRUE_STRINGS


def _normalize_name(s):
    return ' '.join(str(s).strip().lower().split())


def load_tracker_file(path):
    if not path:
        return None
    if not os.path.exists(path):
        logger.warning("Tracker file not found at '%s'.", path)
        return None
    try:
        xl = pd.ExcelFile(path, engine='openpyxl')
        if TRACKER_SHEET_NAME:
            if TRACKER_SHEET_NAME not in xl.sheet_names:
                logger.warning(
                    "TRACKER_SHEET_NAME '%s' not found in '%s' (available sheets: %s) — "
                    "skipping tracker this run.", TRACKER_SHEET_NAME, path, xl.sheet_names
                )
                return None
            sheet_name = TRACKER_SHEET_NAME
        else:
            sheet_name = xl.sheet_names[0]
        df = pd.read_excel(path, engine='openpyxl', sheet_name=sheet_name,
                            header=TRACKER_HEADER_ROW - 1)
    except Exception as e:
        logger.warning("Could not read tracker file '%s' (%s) — skipping it this run.", path, e)
        return None

    name_col      = _find_col(df.columns, _RULE_NAME_COL_ALIASES)
    triggered_col = _find_col(df.columns, _TRIGGERED_COL_ALIASES)

    if triggered_col is None and not df.empty:
        logger.warning(
            "Tracker sheet '%s' has no recognizable 'Triggered' column (looked for: %s).",
            sheet_name, _TRIGGERED_COL_ALIASES
        )
    if name_col is None and not df.empty:
        logger.warning(
            "Tracker sheet '%s' has no recognizable Rule Name column — cross-referencing against "
            "QRadar and auto-mark-triggered logic are disabled this run.", sheet_name
        )

    return {
        'raw':                       df,
        'sheet_name':                sheet_name,
        'name_col':                  name_col,
        'triggered_col':             triggered_col,
        'changed_cells':             [],
        'auto_marked_rows':          [],
        'auto_marked_indices':       set(),
        'unmatched_count':           None,
        'under_investigation_count': 0,
        'triggered_count':           0,
    }


def reconcile_tracker_with_qradar(tracker_info, master_df):
    df            = tracker_info['raw']
    name_col      = tracker_info['name_col']
    triggered_col = tracker_info['triggered_col']

    lookup = {}
    if not master_df.empty:
        for _, r in master_df.iterrows():
            key = _normalize_name(r['rule_name'])
            if key not in lookup:
                lookup[key] = r

    matched_flags, recently_active_list, total_triggered_list = [], [], []
    changed_cells, auto_marked_rows, auto_marked_indices = [], [], set()

    for idx, row in df.iterrows():
        matched_rule = lookup.get(_normalize_name(row[name_col])) if name_col else None

        is_triggered = _to_bool(row[triggered_col]) if triggered_col else False
        recently_active = bool(matched_rule['is_recently_triggered']) if matched_rule is not None else None
        total_triggered = int(matched_rule['triggered_count']) if matched_rule is not None else None

        matched_flags.append(matched_rule is not None)
        recently_active_list.append(recently_active)
        total_triggered_list.append(total_triggered)

        should_auto_mark = (
            AUTO_UPDATE_TRACKER_TRIGGERED
            and matched_rule is not None
            and not is_triggered
            and triggered_col is not None
            and recently_active
        )
        if should_auto_mark:
            df.at[idx, triggered_col] = True
            changed_cells.append((idx, triggered_col, True))
            auto_marked_indices.add(idx)
            auto_marked_rows.append({
                'rule_name': str(row[name_col]),
                'last_triggered': matched_rule['last_triggered'],
            })

    df['_matched_to_qradar']    = matched_flags
    df['_recently_active']      = recently_active_list
    df['_triggered_count']      = total_triggered_list

    tracker_info['raw'] = df
    tracker_info['changed_cells'] = changed_cells
    tracker_info['auto_marked_rows'] = auto_marked_rows
    tracker_info['auto_marked_indices'] = auto_marked_indices
    tracker_info['unmatched_count'] = int(len(matched_flags) - sum(matched_flags)) if name_col else None
    tracker_info['triggered_count'] = int(df[triggered_col].apply(_to_bool).sum()) if triggered_col else 0
    tracker_info['under_investigation_count'] = (
        int((~df[triggered_col].apply(_to_bool)).sum()) if triggered_col else 0
    )
    return tracker_info


def persist_tracker_updates(path, tracker_info):
    changed_cells = tracker_info.get('changed_cells') or []
    if not changed_cells:
        return 'no_changes'
    if not AUTO_UPDATE_TRACKER_TRIGGERED:
        return 'disabled'

    try:
        if TRACKER_BACKUP_BEFORE_WRITE:
            backup_path = path.replace('.xlsx', f"_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
            shutil.copy2(path, backup_path)

        wb = openpyxl.load_workbook(path)
        ws = wb[tracker_info['sheet_name']]
        header_cells = ws[TRACKER_HEADER_ROW]
        col_letter_by_name = {str(c.value).strip(): get_column_letter(c.column) for c in header_cells if c.value}

        for idx, col_name, value in changed_cells:
            col_letter = col_letter_by_name.get(col_name)
            if not col_letter:
                continue
            excel_row = idx + TRACKER_HEADER_ROW + 1
            ws[f'{col_letter}{excel_row}'] = value

        wb.save(path)
        return 'written'
    except PermissionError:
        logger.warning("Tracker file '%s' is open elsewhere — could not persist updates this run.", path)
        return 'locked'
    except Exception as e:
        logger.error("Could not write updates back to tracker file '%s': %s", path, e)
        return 'error'


def build_master_rule_table(ucm_rows):
    newly_cutoff_ms  = int((datetime.now() - timedelta(days=NEWLY_TRIGGERED_WINDOW_DAYS)).timestamp() * 1000)
    recent_cutoff_ms = int((datetime.now() - timedelta(days=TRACKER_VALIDATION_WINDOW_DAYS)).timestamp() * 1000)

    rows = []
    for raw in ucm_rows:
        row = _normalize_ucm_row(raw, UCM_REPORT_COLUMNS)

        rule_id  = _row_get(row, 'ID', 'id', 'Rule_ID')
        name     = _row_get(row, 'N', 'name', 'Attribute_Name') or f'Rule {rule_id}'

        created_ts, created_display   = _parse_ucm_date(_row_get(row, 'CD', 'creation_date', 'Attribute_Creation_Date'))
        _,          modified_display  = _parse_ucm_date(_row_get(row, 'MD', 'modification_date', 'Attribute_Modification_Date'))
        first_ts,   first_display     = _parse_ucm_date(_row_get(row, 'FTG', 'first_triggered', 'Attribute_First_Triggered'))
        last_ts,    last_display      = _parse_ucm_date(_row_get(row, 'LTG', 'last_triggered', 'Attribute_Last_Triggered'))
        triggered_count = _parse_ucm_triggered_count(_row_get(row, 'TG', 'triggered', 'Attribute_Triggered'))

        rows.append({
            'rule_id':               rule_id,
            'rule_name':             _csv_formula_safe(name),
            'rule_type':             _row_get(row, 'T', 'type', 'Attribute_Type') or 'UNKNOWN',
            'origin':                _row_get(row, 'RO', 'origin', 'Rule_Origin') or 'UNKNOWN',
            'group':                 _row_get(row, 'GR', 'group', 'Attribute_Group') or '',
            'action':                _row_get(row, 'A', 'action', 'Attribute_Action') or '',
            'response':              _row_get(row, 'RE', 'response', 'Attribute_Response') or '',
            'action_details':        _row_get(row, 'AD', 'action_details', 'Action_Details') or '',
            'response_details':      _row_get(row, 'RED', 'response_details', 'Response_Details') or '',
            'created':               created_display,
            'created_age_days':      _age_days_from_epoch(created_ts) if created_ts else None,
            'modified':              modified_display,
            'description':           _csv_formula_safe(
                                          _row_get(row, 'NO', 'notes', 'Attribute_Notes')
                                          or 'Not returned by Use Case Explorer for this rule.'
                                      ),
            'triggered_count':       triggered_count,
            'first_triggered':       first_display,
            'first_triggered_ts':    first_ts,
            'last_triggered':        last_display,
            'last_triggered_ts':     last_ts,
            'is_dead':               (first_ts is None) and triggered_count == 0,
            'is_newly_triggered':    bool(first_ts and first_ts >= newly_cutoff_ms),
            'is_recently_triggered': bool(last_ts and last_ts >= recent_cutoff_ms),
        })
    return pd.DataFrame(rows)


_FILLS = {
    'red':         PatternFill(start_color='FF6B6B', end_color='FF6B6B', fill_type='solid'),
    'orange':      PatternFill(start_color='FFBF47', end_color='FFBF47', fill_type='solid'),
    'green':       PatternFill(start_color='A8E6CF', end_color='A8E6CF', fill_type='solid'),
    'blue':        PatternFill(start_color='74B9FF', end_color='74B9FF', fill_type='solid'),
    'teal':        PatternFill(start_color='7FE0D0', end_color='7FE0D0', fill_type='solid'),
    'header':      PatternFill(start_color='2D2257', end_color='2D2257', fill_type='solid'),
    'zebra_a':     PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid'),
    'zebra_b':     PatternFill(start_color='F3F1FB', end_color='F3F1FB', fill_type='solid'),
    'stale':       PatternFill(start_color='FFE2E2', end_color='FFE2E2', fill_type='solid'),
    'auto_marked': PatternFill(start_color='C6F6D5', end_color='C6F6D5', fill_type='solid'),
}
_HDR_FONT   = Font(bold=True, color='E8E0FF', size=10)
_BOLD       = Font(bold=True, size=10)
_STALE_FONT = Font(bold=True, size=10, color='B91C1C')
_CENTRE     = Alignment(horizontal='center', vertical='center')
_WRAP       = Alignment(horizontal='left',   vertical='top', wrap_text=True)
_THIN_SIDE  = Side(style='thin', color='DCD7EE')
_THIN_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)


def _write_sheet_header(ws, columns, col_widths):
    ws.append(columns)
    for col_idx, (cell, width) in enumerate(zip(ws[1], col_widths), start=1):
        cell.fill      = _FILLS['header']
        cell.font      = _HDR_FONT
        cell.alignment = _CENTRE
        cell.border    = _THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = 'A2'


def _apply_zebra_and_borders(ws, first_data_row, last_data_row, num_cols):
    for r in range(first_data_row, last_data_row + 1):
        fill = _FILLS['zebra_a'] if (r - first_data_row) % 2 == 0 else _FILLS['zebra_b']
        for c in range(1, num_cols + 1):
            cell = ws.cell(row=r, column=c)
            if cell.fill.start_color.rgb in (None, '00000000'):
                cell.fill = fill
            cell.border = _THIN_BORDER


def _pct(n, total):
    return f"{(n / total * 100):.1f}%" if total else "0.0%"


def _write_summary_sheet(wb, dead_df, newly_df, tracker_info, total_rules):
    ws = wb.active
    ws.title = 'Executive Summary'

    ws['A1'] = f'QRadar — {REPORT_TITLE}'
    ws['A1'].font = Font(bold=True, size=16, color='2D2257')
    ws['A2'] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws['A2'].font = Font(italic=True, size=10, color='7C6FA0')

    ws['A3'] = (f"Data source: QRadar Use Case Manager (Use Case Explorer)   ·   "
                f"Newly-triggered window: last {NEWLY_TRIGGERED_WINDOW_DAYS} day(s)   ·   "
                f"Tracker validation window: last {TRACKER_VALIDATION_WINDOW_DAYS} day(s)")
    ws['A3'].font = Font(size=10, color='7C6FA0')

    ws['A5'] = 'Rules matching filters (enabled, not a building block, offense-creating action):'
    ws['B5'] = int(total_rules)
    ws['B5'].font = _BOLD

    row = 7
    ws.cell(row=row, column=1, value='Metric').fill = _FILLS['header']
    ws.cell(row=row, column=1).font = _HDR_FONT
    for c, label in ((2, 'Count'), (3, '% of Rules')):
        ws.cell(row=row, column=c, value=label).fill = _FILLS['header']
        ws.cell(row=row, column=c).font = _HDR_FONT
        ws.cell(row=row, column=c).alignment = _CENTRE
    row += 1

    dead_n  = len(dead_df)
    newly_n = len(newly_df)
    metrics = [
        ('Dead — Never Triggered',  dead_n,  _pct(dead_n, total_rules),  'red'),
        (f'Newly Triggered (≤{NEWLY_TRIGGERED_WINDOW_DAYS}d)', newly_n, _pct(newly_n, total_rules), 'teal'),
        ('Under Investigation',     tracker_info['under_investigation_count'] if tracker_info else 'N/A', '', 'orange'),
        ('Triggered (confirmed)',   tracker_info['triggered_count']           if tracker_info else 'N/A', '', 'blue'),
    ]
    for label, value, pct, fill_key in metrics:
        ws.cell(row=row, column=1, value=label).font = _BOLD
        c1 = ws.cell(row=row, column=2, value=value)
        c1.alignment = _CENTRE
        if value != 'N/A':
            c1.fill = _FILLS[fill_key]
        ws.cell(row=row, column=3, value=pct).alignment = _CENTRE
        row += 1
    row += 1

    if tracker_info is None:
        tracker_note = (
            f"Tracker not configured or not found (TRACKER_EXCEL_PATH = '{TRACKER_EXCEL_PATH or '(blank)'}')."
        )
    else:
        crossref = 'enabled (matched by Rule Name)' if tracker_info['name_col'] else 'disabled (no Rule Name column found)'
        auto_n = len(tracker_info['auto_marked_rows'])
        unmatched = tracker_info['unmatched_count']
        tracker_note = (
            f"Tracker source: {TRACKER_EXCEL_PATH}  (sheet: {tracker_info['sheet_name']})  ·  "
            f"{len(tracker_info['raw'])} row(s) read  ·  Cross-referencing: {crossref}  ·  "
            f"Auto-marked Triggered this run: {auto_n}"
            + (f"  ·  Unmatched rows: {unmatched}" if unmatched is not None else "")
        )
    ws.cell(row=row, column=1, value=tracker_note).font = Font(italic=True, size=9, color='7C6FA0')
    ws.cell(row=row, column=1).alignment = _WRAP
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.row_dimensions[row].height = 40
    row += 2

    caveat = (
        'All rule data — including First/Last Triggered — comes from QRadar Use Case Manager\'s '
        'Use Case Explorer report, not from raw offense data, so none of this is bounded by '
        'offense retention. "Dead" means never triggered per UCM\'s own stats. If UCM resets its '
        'triggered stats when a rule is edited/re-saved in your deployment, a recently-modified '
        'rule can show as never-triggered even though it fired before its last edit — worth '
        'spot-checking a known-good rule if a "Dead" count looks surprising. Auto-marking '
        'Triggered uses recent trigger activity as a proxy for "fired since you started testing" '
        '— it does not confirm the trigger was caused by your specific test versus real traffic.'
    )
    ws.cell(row=row, column=1, value=caveat).font = Font(italic=True, size=9, color='7C6FA0')
    ws.cell(row=row, column=1).alignment = _WRAP
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.row_dimensions[row].height = 75

    for i, w in enumerate([46, 14, 14, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_dead_sheet(wb, dead_df):
    ws = wb.create_sheet('Dead Rules')
    cols   = ['Rule ID', 'Rule Name', 'Type', 'Origin', 'Group', 'Created', 'Age (Days)',
              'Modified', 'Response', 'Description / Notes']
    widths = [9, 42, 14, 10, 16, 12, 11, 12, 24, 55]
    _write_sheet_header(ws, cols, widths)

    if dead_df.empty:
        ws.append(['No dead rules — every in-scope rule has triggered at least once.'])
        return

    sorted_df = dead_df.sort_values('created_age_days', ascending=False, na_position='last')
    desc_idx = len(cols)
    age_idx  = cols.index('Age (Days)') + 1
    first_row = ws.max_row + 1

    for _, r in sorted_df.iterrows():
        age = r['created_age_days']
        ws.append([int(r['rule_id']) if r['rule_id'] is not None else '—', str(r['rule_name']),
                   str(r['rule_type']), str(r['origin']), str(r['group']), str(r['created']),
                   age if age is not None else '—', str(r['modified']), str(r['response']),
                   str(r['description'])])
        r_idx = ws.max_row
        ws.cell(row=r_idx, column=desc_idx).alignment = _WRAP
        if age is not None and age >= DEAD_RULE_STALE_AGE_DAYS:
            ws.cell(row=r_idx, column=age_idx).fill = _FILLS['stale']
            ws.cell(row=r_idx, column=age_idx).font = _STALE_FONT

    _apply_zebra_and_borders(ws, first_row, ws.max_row, len(cols))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"


def _write_newly_triggered_sheet(wb, newly_df):
    ws = wb.create_sheet('Newly Triggered Rules')
    cols   = ['Rule ID', 'Rule Name', 'Type', 'Group', 'First Triggered', 'Last Triggered', 'Triggered Count']
    widths = [9, 42, 14, 16, 15, 15, 16]
    _write_sheet_header(ws, cols, widths)

    if newly_df.empty:
        ws.append([f'No rules first triggered in the last {NEWLY_TRIGGERED_WINDOW_DAYS} day(s).'])
        return

    sorted_df = newly_df.sort_values('first_triggered_ts', ascending=False, na_position='last')
    first_row = ws.max_row + 1

    for _, r in sorted_df.iterrows():
        ws.append([int(r['rule_id']) if r['rule_id'] is not None else '—', str(r['rule_name']),
                   str(r['rule_type']), str(r['group']), str(r['first_triggered']),
                   str(r['last_triggered']), int(r['triggered_count'])])

    _apply_zebra_and_borders(ws, first_row, ws.max_row, len(cols))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"


def _write_tracker_sheet(wb, tracker_info):
    ws = wb.create_sheet('Rule Testing Tracker')
    df = tracker_info['raw']
    original_cols = [c for c in df.columns if not str(c).startswith('_')]
    has_match_info = '_matched_to_qradar' in df.columns
    extra_cols = (
        ['Matched to QRadar Rule?', f'Fired in last {TRACKER_VALIDATION_WINDOW_DAYS}d?',
         'Triggered Count (Total)', 'Auto-Marked Triggered This Run?']
        if has_match_info else []
    )
    cols = [str(c) for c in original_cols] + extra_cols
    widths = [max(14, min(50, len(str(c)) + 4)) for c in original_cols] + [20, 20, 18, 22][:len(extra_cols)]
    _write_sheet_header(ws, cols, widths)

    if df.empty:
        ws.append(['Tracker sheet has no rows.'])
        return

    triggered_col = tracker_info['triggered_col']
    auto_marked_indices = tracker_info.get('auto_marked_indices', set())
    first_row = ws.max_row + 1

    for idx, r in df.iterrows():
        values = [_native(r[c]) for c in original_cols]
        if has_match_info:
            active, total = r['_recently_active'], r['_triggered_count']
            values += [
                'Yes' if r['_matched_to_qradar'] else 'No',
                ('Yes' if active else 'No') if active is not None else '—',
                int(total) if pd.notna(total) else '—',
                'Yes' if idx in auto_marked_indices else 'No',
            ]
        ws.append(values)
        r_idx = ws.max_row

        if idx in auto_marked_indices:
            for c_idx in range(1, len(cols) + 1):
                ws.cell(row=r_idx, column=c_idx).fill = _FILLS['auto_marked']

        for c_idx, c in enumerate(original_cols, start=1):
            if c == triggered_col:
                if _to_bool(r[c]):
                    ws.cell(row=r_idx, column=c_idx).fill = _FILLS['blue']
                else:
                    ws.cell(row=r_idx, column=c_idx).fill = _FILLS['orange']

    _apply_zebra_and_borders(ws, first_row, ws.max_row, len(cols))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"


def save_excel_report(dead_df, newly_df, tracker_info, total_rules, path):
    try:
        wb = openpyxl.Workbook()
        _write_summary_sheet(wb, dead_df, newly_df, tracker_info, total_rules)
        _write_dead_sheet(wb, dead_df)
        _write_newly_triggered_sheet(wb, newly_df)
        if tracker_info is not None:
            _write_tracker_sheet(wb, tracker_info)

        try:
            wb.save(path)
            print(f"✅ Excel saved → {path}")
            return path
        except PermissionError:
            fallback = path.replace('.xlsx', f"_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
            print(f"⚠  '{path}' is open elsewhere — saving to '{fallback}' instead.")
            wb.save(fallback)
            print(f"✅ Excel saved → {fallback}")
            return fallback
    except Exception as e:
        logger.error("Excel save failed:\n%s", traceback.format_exc())
        print(f"❌ Excel save error: {e}")
        return None


_CHART_BG    = '#0a0a10'
_CHART_GRID  = '#201c33'
_CHART_TEXT  = '#cfc8ea'
_CHART_TITLE = '#a78bfa'
_CHART_AXIS  = '#332c50'


def generate_status_chart(dead_count, newly_count, tracker_info, total_rules):
    labels  = ['Dead', 'Newly\nTriggered']
    values  = [dead_count, newly_count]
    colours = ['#f87171', '#2dd4bf']

    if tracker_info is not None:
        labels  += ['Under\nInvestigation', 'Triggered']
        values  += [tracker_info['under_investigation_count'], tracker_info['triggered_count']]
        colours += ['#f5b155', '#5b8cff']

    fig, ax = plt.subplots(figsize=(6.4, 4.0), facecolor=_CHART_BG)
    ax.set_facecolor(_CHART_BG)
    ax.yaxis.grid(True, color=_CHART_GRID, linewidth=0.8, zorder=0)
    ax.set_axisbelow(True)

    bars = ax.bar(labels, values, color=colours, edgecolor=_CHART_BG, linewidth=1.2, zorder=3, width=0.5)
    max_val = max(values + [1])
    for bar, val in zip(bars, values):
        pct = f"{val / total_rules * 100:.1f}% of rules" if total_rules else ""
        ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + max_val * 0.03, str(val),
                ha='center', va='bottom', color=_CHART_TEXT, fontsize=11, fontweight='bold')
        if pct:
            ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() * 0.5, pct,
                    ha='center', va='center', color=_CHART_BG, fontsize=7.5)

    ax.set_ylim(0, max_val * 1.25)
    ax.set_title('Rule Status Overview', color=_CHART_TITLE, fontsize=12, fontweight='700', pad=14)
    ax.tick_params(colors=_CHART_TEXT, labelsize=9.5)
    for spine in ax.spines.values():
        spine.set_edgecolor(_CHART_AXIS)
        spine.set_linewidth(0.6)
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    plt.tight_layout(pad=1.0)

    tmp  = tempfile.NamedTemporaryFile(suffix='.png', delete=False, prefix='qr_status_')
    path = tmp.name
    tmp.close()
    plt.savefig(path, bbox_inches='tight', dpi=120, facecolor=_CHART_BG, edgecolor='none')
    plt.close(fig)
    return path


_C = {
    'page':   '#0a0a10',
    'card':   '#111017',
    'border': '#26223a',
    'purple': '#8b5cf6',
    'violet': '#a78bfa',
    'dim':    '#8d87a8',
    'text':   '#f4f2fb',
    'red':    '#f87171',
    'amber':  '#f5b155',
    'blue':   '#5b8cff',
    'teal':   '#2dd4bf',
}


def _dot(color):
    return f'<span style="color:{color};font-size:10px;vertical-align:middle;">●</span>'


def _build_dead_table_html(dead_df):
    C = _C
    if dead_df.empty:
        return (f'<p style="color:{C["blue"]};font-size:12px;font-weight:600;padding:8px 0;margin:0;">'
                f'{_dot(C["blue"])} No dead rules — every in-scope rule has triggered at least once.</p>')

    shown = dead_df.sort_values('created_age_days', ascending=False, na_position='last').head(EMAIL_TABLE_ROW_CAP)
    rows_html = ''
    for i, (_, row) in enumerate(shown.iterrows()):
        bg = C['page'] if i % 2 == 0 else C['card']
        name = _html_escape(row['rule_name'])
        name_short = name[:42] + '…' if len(name) > 42 else name
        age = row['created_age_days']
        age_txt = f"{age}d" if age is not None else '—'
        rows_html += f"""
        <tr bgcolor="{bg}" style="background-color:{bg};">
          <td bgcolor="{bg}" style="padding:9px 12px;border-left:2px solid {C['red']};font-size:12px;color:{C['text']};" title="{name}">{name_short}</td>
          <td bgcolor="{bg}" style="padding:9px 12px;font-size:11px;color:{C['dim']};text-align:center;">{_html_escape(row['rule_type'])}</td>
          <td bgcolor="{bg}" style="padding:9px 12px;font-size:11px;color:{C['dim']};text-align:center;">{_html_escape(row['group'])}</td>
          <td bgcolor="{bg}" style="padding:9px 12px;font-size:11px;color:{C['dim']};text-align:center;">{_html_escape(row['created'])}</td>
          <td bgcolor="{bg}" style="padding:9px 12px;font-size:11px;color:{C['red']};text-align:center;font-weight:600;">{age_txt}</td>
        </tr>"""

    more_note = ''
    if len(dead_df) > EMAIL_TABLE_ROW_CAP:
        more_note = (f'<div style="font-size:10px;color:{C["dim"]};margin-top:8px;">'
                     f'+ {len(dead_df) - EMAIL_TABLE_ROW_CAP} more — see attached Excel for the full list.</div>')

    return f"""
    <table width="100%" cellpadding="0" cellspacing="0" bgcolor="{C['page']}" style="border-collapse:collapse;margin-top:10px;background-color:{C['page']};">
      <thead><tr>
        <th bgcolor="{C['page']}" style="padding:8px 12px;text-align:left;font-size:9px;color:{C['red']};font-weight:700;text-transform:uppercase;letter-spacing:1px;border-bottom:1px solid {C['border']};">Rule Name</th>
        <th bgcolor="{C['page']}" style="padding:8px 12px;text-align:center;font-size:9px;color:{C['red']};font-weight:700;text-transform:uppercase;letter-spacing:1px;border-bottom:1px solid {C['border']};">Type</th>
        <th bgcolor="{C['page']}" style="padding:8px 12px;text-align:center;font-size:9px;color:{C['red']};font-weight:700;text-transform:uppercase;letter-spacing:1px;border-bottom:1px solid {C['border']};">Group</th>
        <th bgcolor="{C['page']}" style="padding:8px 12px;text-align:center;font-size:9px;color:{C['red']};font-weight:700;text-transform:uppercase;letter-spacing:1px;border-bottom:1px solid {C['border']};">Created</th>
        <th bgcolor="{C['page']}" style="padding:8px 12px;text-align:center;font-size:9px;color:{C['red']};font-weight:700;text-transform:uppercase;letter-spacing:1px;border-bottom:1px solid {C['border']};">Age</th>
      </tr></thead>
      <tbody>{rows_html}</tbody>
    </table>{more_note}"""


def _build_newly_triggered_table_html(newly_df):
    C = _C
    if newly_df.empty:
        return (f'<p style="color:{C["dim"]};font-size:12px;padding:8px 0;margin:0;">'
                f'No rules first triggered in the last {NEWLY_TRIGGERED_WINDOW_DAYS} day(s).</p>')

    shown = newly_df.sort_values('first_triggered_ts', ascending=False, na_position='last').head(EMAIL_TABLE_ROW_CAP)
    rows_html = ''
    for i, (_, row) in enumerate(shown.iterrows()):
        bg = C['page'] if i % 2 == 0 else C['card']
        name = _html_escape(row['rule_name'])
        name_short = name[:42] + '…' if len(name) > 42 else name
        rows_html += f"""
        <tr bgcolor="{bg}" style="background-color:{bg};">
          <td bgcolor="{bg}" style="padding:9px 12px;border-left:2px solid {C['teal']};font-size:12px;color:{C['text']};" title="{name}">{name_short}</td>
          <td bgcolor="{bg}" style="padding:9px 12px;font-size:11px;color:{C['dim']};text-align:center;">{_html_escape(row['group'])}</td>
          <td bgcolor="{bg}" style="padding:9px 12px;font-size:11px;color:{C['teal']};text-align:center;font-weight:600;">{_html_escape(row['first_triggered'])}</td>
          <td bgcolor="{bg}" style="padding:9px 12px;font-size:11px;color:{C['dim']};text-align:center;">{int(row['triggered_count'])}</td>
        </tr>"""

    more_note = ''
    if len(newly_df) > EMAIL_TABLE_ROW_CAP:
        more_note = (f'<div style="font-size:10px;color:{C["dim"]};margin-top:8px;">'
                     f'+ {len(newly_df) - EMAIL_TABLE_ROW_CAP} more — see attached Excel for the full list.</div>')

    return f"""
    <table width="100%" cellpadding="0" cellspacing="0" bgcolor="{C['page']}" style="border-collapse:collapse;margin-top:10px;background-color:{C['page']};">
      <thead><tr>
        <th bgcolor="{C['page']}" style="padding:8px 12px;text-align:left;font-size:9px;color:{C['teal']};font-weight:700;text-transform:uppercase;letter-spacing:1px;border-bottom:1px solid {C['border']};">Rule Name</th>
        <th bgcolor="{C['page']}" style="padding:8px 12px;text-align:center;font-size:9px;color:{C['teal']};font-weight:700;text-transform:uppercase;letter-spacing:1px;border-bottom:1px solid {C['border']};">Group</th>
        <th bgcolor="{C['page']}" style="padding:8px 12px;text-align:center;font-size:9px;color:{C['teal']};font-weight:700;text-transform:uppercase;letter-spacing:1px;border-bottom:1px solid {C['border']};">First Triggered</th>
        <th bgcolor="{C['page']}" style="padding:8px 12px;text-align:center;font-size:9px;color:{C['teal']};font-weight:700;text-transform:uppercase;letter-spacing:1px;border-bottom:1px solid {C['border']};">Triggered Count</th>
      </tr></thead>
      <tbody>{rows_html}</tbody>
    </table>{more_note}"""


def _build_under_investigation_table_html(tracker_info):
    C = _C
    if tracker_info is None or not tracker_info.get('triggered_col'):
        return (f'<p style="color:{C["dim"]};font-size:12px;padding:8px 0;margin:0;">'
                f'Tracker not configured, or no Triggered column found.</p>')

    df = tracker_info['raw']
    triggered_col, name_col = tracker_info['triggered_col'], tracker_info['name_col']
    has_match = '_matched_to_qradar' in df.columns

    mask = ~df[triggered_col].apply(_to_bool)
    shown = df[mask]
    if shown.empty:
        return (f'<p style="color:{C["blue"]};font-size:12px;font-weight:600;padding:8px 0;margin:0;">'
                f'{_dot(C["blue"])} Nothing currently under investigation — everything tracked is Triggered.</p>')

    shown = shown.head(EMAIL_TABLE_ROW_CAP)
    rows_html = ''
    for i, (_, row) in enumerate(shown.iterrows()):
        bg = C['page'] if i % 2 == 0 else C['card']
        name = _html_escape(row[name_col]) if name_col else '(no name column)'
        name_short = name[:42] + '…' if len(name) > 42 else name
        active = row['_recently_active'] if has_match else None
        if active is not None:
            fired_txt, fired_color = ('Yes', C['blue']) if active else ('No', C['dim'])
        else:
            fired_txt, fired_color = '—', C['dim']
        rows_html += f"""
        <tr bgcolor="{bg}" style="background-color:{bg};">
          <td bgcolor="{bg}" style="padding:9px 12px;border-left:2px solid {C['amber']};font-size:12px;color:{C['text']};" title="{name}">{name_short}</td>
          <td bgcolor="{bg}" style="padding:9px 12px;font-size:11px;color:{fired_color};text-align:center;font-weight:600;">{fired_txt}</td>
        </tr>"""

    more_note = ''
    if mask.sum() > EMAIL_TABLE_ROW_CAP:
        more_note = (f'<div style="font-size:10px;color:{C["dim"]};margin-top:8px;">'
                     f'+ {int(mask.sum()) - EMAIL_TABLE_ROW_CAP} more — see attached Excel for the full list.</div>')

    return f"""
    <table width="100%" cellpadding="0" cellspacing="0" bgcolor="{C['page']}" style="border-collapse:collapse;margin-top:10px;background-color:{C['page']};">
      <thead><tr>
        <th bgcolor="{C['page']}" style="padding:8px 12px;text-align:left;font-size:9px;color:{C['amber']};font-weight:700;text-transform:uppercase;letter-spacing:1px;border-bottom:1px solid {C['border']};">Rule Name</th>
        <th bgcolor="{C['page']}" style="padding:8px 12px;text-align:center;font-size:9px;color:{C['amber']};font-weight:700;text-transform:uppercase;letter-spacing:1px;border-bottom:1px solid {C['border']};">Fired in last {TRACKER_VALIDATION_WINDOW_DAYS}d?</th>
      </tr></thead>
      <tbody>{rows_html}</tbody>
    </table>{more_note}"""


def _build_auto_validated_table_html(tracker_info):
    C = _C
    rows = tracker_info.get('auto_marked_rows') if tracker_info else None
    if not rows:
        return None

    shown = rows[:EMAIL_TABLE_ROW_CAP]
    rows_html = ''
    for i, r in enumerate(shown):
        bg = C['page'] if i % 2 == 0 else C['card']
        name = _html_escape(r['rule_name'])
        name_short = name[:48] + '…' if len(name) > 48 else name
        rows_html += f"""
        <tr bgcolor="{bg}" style="background-color:{bg};">
          <td bgcolor="{bg}" style="padding:9px 12px;border-left:2px solid {C['blue']};font-size:12px;color:{C['text']};" title="{name}">{name_short}</td>
          <td bgcolor="{bg}" style="padding:9px 12px;font-size:11px;color:{C['blue']};text-align:center;font-weight:600;">{_html_escape(r['last_triggered'])}</td>
        </tr>"""

    more_note = ''
    if len(rows) > EMAIL_TABLE_ROW_CAP:
        more_note = (f'<div style="font-size:10px;color:{C["dim"]};margin-top:8px;">'
                     f'+ {len(rows) - EMAIL_TABLE_ROW_CAP} more — see attached Excel.</div>')

    return f"""
    <table width="100%" cellpadding="0" cellspacing="0" bgcolor="{C['page']}" style="border-collapse:collapse;margin-top:10px;background-color:{C['page']};">
      <thead><tr>
        <th bgcolor="{C['page']}" style="padding:8px 12px;text-align:left;font-size:9px;color:{C['blue']};font-weight:700;text-transform:uppercase;letter-spacing:1px;border-bottom:1px solid {C['border']};">Rule Name</th>
        <th bgcolor="{C['page']}" style="padding:8px 12px;text-align:center;font-size:9px;color:{C['blue']};font-weight:700;text-transform:uppercase;letter-spacing:1px;border-bottom:1px solid {C['border']};">Last Triggered</th>
      </tr></thead>
      <tbody>{rows_html}</tbody>
    </table>{more_note}"""


def build_email_html(dead_df, newly_df, tracker_info, total_rules, chart_cid):
    C = _C
    run_time = datetime.now().strftime('%d %b %Y  ·  %H:%M:%S')

    dead_count   = len(dead_df)
    newly_count  = len(newly_df)
    inv_count    = tracker_info['under_investigation_count'] if tracker_info else None
    trig_count   = tracker_info['triggered_count']            if tracker_info else None
    auto_marked  = tracker_info['auto_marked_rows']            if tracker_info else []

    if dead_count > 0:
        hdr_txt, hdr_color = f'{dead_count} DEAD RULE(S) TO REVIEW', C['red']
    elif newly_count > 0:
        hdr_txt, hdr_color = f'{newly_count} NEWLY TRIGGERED RULE(S)', C['teal']
    elif auto_marked:
        hdr_txt, hdr_color = f'{len(auto_marked)} AUTO-VALIDATED THIS RUN', C['blue']
    else:
        hdr_txt, hdr_color = 'STEADY STATE', C['violet']

    def badge(color, txt):
        return (f'<span style="color:{color};font-size:10px;font-weight:700;'
                f'padding:5px 13px;border-radius:20px;letter-spacing:0.6px;border:1px solid {color}80;'
                f'white-space:nowrap;">{txt}</span>')

    def metric_card(label, value, color, note=''):
        display_val = value if value is not None else '—'
        note_html = (f'<div style="font-size:9px;color:{C["dim"]};margin-top:4px;">{note}</div>'
                     if note else '')
        return (f'<td width="25%" bgcolor="{C["page"]}" style="padding:5px;background-color:{C["page"]};">'
                f'<table width="100%" cellpadding="0" cellspacing="0" bgcolor="{C["card"]}" '
                f'style="border-collapse:collapse;background-color:{C["card"]};border:1px solid {C["border"]};'
                f'border-top:2px solid {color};border-radius:8px;">'
                f'<tr><td bgcolor="{C["card"]}" style="background-color:{C["card"]};padding:14px 6px;text-align:center;">'
                f'<div style="font-size:24px;font-weight:800;color:{color};line-height:1;letter-spacing:-1px;">{display_val}</div>'
                f'<div style="font-size:8.5px;color:{C["dim"]};margin-top:6px;text-transform:uppercase;letter-spacing:0.8px;">{label}</div>'
                f'{note_html}</td></tr></table></td>')

    dead_pct  = f"{dead_count / total_rules * 100:.1f}% of rules" if total_rules else ''
    newly_pct = f"{newly_count / total_rules * 100:.1f}% of rules" if total_rules else ''
    inv_note = '' if tracker_info else 'not configured'
    trig_note = f"+{len(auto_marked)} auto this run" if auto_marked else ('' if tracker_info else 'not configured')

    headline_metrics = (
        metric_card('Dead Rules', dead_count, C['red'], dead_pct)
        + metric_card(f'Newly Triggered ({NEWLY_TRIGGERED_WINDOW_DAYS}d)', newly_count, C['teal'], newly_pct)
        + metric_card('Under Investigation', inv_count, C['amber'], inv_note)
        + metric_card('Triggered', trig_count, C['blue'], trig_note)
    )

    chart_html = (f'<img src="cid:{chart_cid}" alt="Rule status chart" '
                  f'style="display:block;max-width:100%;margin:18px auto 0;">') if chart_cid else ''
    dead_table_html   = _build_dead_table_html(dead_df)
    newly_table_html  = _build_newly_triggered_table_html(newly_df)
    inv_table_html    = _build_under_investigation_table_html(tracker_info)
    auto_table_html   = _build_auto_validated_table_html(tracker_info)

    def section_header(color, title, subtitle):
        return f"""
  <tr><td bgcolor="{C['page']}" style="background-color:{C['page']};padding:28px 0 4px;border-top:1px solid {C['border']};">
    <span style="font-size:13px;font-weight:700;color:{C['text']};">{_dot(color)}&nbsp;&nbsp;{title}</span>
    <div style="font-size:10px;color:{C['dim']};margin-top:4px;">{subtitle}</div>
  </td></tr>"""

    auto_section = ''
    if auto_table_html:
        auto_section = (
            section_header(C['blue'], 'Auto-Validated This Run',
                            'Was marked not-yet-triggered but has fired since — Triggered has been set to Yes in your tracker file.')
            + f'<tr><td bgcolor="{C["page"]}" style="background-color:{C["page"]};padding:0 0 20px;">{auto_table_html}</td></tr>'
        )

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<meta name="color-scheme" content="dark">
<meta name="supported-color-schemes" content="dark">
<title>{REPORT_TITLE}</title>
<!--[if mso]>
<style type="text/css">
table, td {{ font-family:Segoe UI, Arial, sans-serif !important; }}
</style>
<![endif]-->
</head>
<body bgcolor="{C['page']}" style="margin:0;padding:0;background-color:{C['page']};font-family:'Segoe UI',Helvetica,Arial,sans-serif;">
<table role="presentation" width="100%" bgcolor="{C['page']}" cellpadding="0" cellspacing="0" style="background-color:{C['page']};">
<tr>
<td align="center" bgcolor="{C['page']}" style="background-color:{C['page']};padding:32px 16px;">

<table role="presentation" width="640" cellpadding="0" cellspacing="0" bgcolor="{C['page']}" style="width:640px;max-width:640px;background-color:{C['page']};">

  <tr><td bgcolor="{C['page']}" style="background-color:{C['page']};padding:0 0 18px;border-bottom:2px solid {C['purple']};">
    <table width="100%" cellpadding="0" cellspacing="0" bgcolor="{C['page']}" style="background-color:{C['page']};"><tr>
      <td bgcolor="{C['page']}" style="background-color:{C['page']};">
        <div style="font-size:9px;color:{C['dim']};letter-spacing:3px;text-transform:uppercase;margin-bottom:8px;">
          QRadar Use Case Manager
        </div>
        <div style="font-size:23px;font-weight:800;color:{C['text']};letter-spacing:-0.4px;line-height:1.25;">
          {REPORT_TITLE}
        </div>
        <div style="margin-top:8px;font-size:11px;color:{C['dim']};">{run_time}</div>
      </td>
      <td align="right" valign="top" bgcolor="{C['page']}" style="background-color:{C['page']};">{badge(hdr_color, hdr_txt)}</td>
    </tr></table>
  </td></tr>

  <tr><td bgcolor="{C['page']}" style="background-color:{C['page']};padding:20px 0 4px;">
    <table width="100%" cellpadding="0" cellspacing="0" bgcolor="{C['page']}" style="background-color:{C['page']};"><tr>{headline_metrics}</tr></table>
  </td></tr>

  <tr><td bgcolor="{C['page']}" style="background-color:{C['page']};padding:0 0 4px;text-align:center;">{chart_html}</td></tr>

  {section_header(C['teal'], 'Newly Triggered Rules',
                   f"Fired for the first time within the last {NEWLY_TRIGGERED_WINDOW_DAYS} day(s) — freshly validated rules.")}
  <tr><td bgcolor="{C['page']}" style="background-color:{C['page']};padding:0 0 8px;">{newly_table_html}</td></tr>

  {section_header(C['red'], 'Dead Rules — Never Triggered',
                   "Per QRadar Use Case Manager. Candidates for review or removal.")}
  <tr><td bgcolor="{C['page']}" style="background-color:{C['page']};padding:0 0 8px;">{dead_table_html}</td></tr>

  {section_header(C['amber'], 'Rules Under Investigation',
                   f"Tracker rows where Triggered is not yet Yes — checked against the last {TRACKER_VALIDATION_WINDOW_DAYS} day(s).")}
  <tr><td bgcolor="{C['page']}" style="background-color:{C['page']};padding:0 0 8px;">{inv_table_html}</td></tr>

  {auto_section}

  <tr><td bgcolor="{C['page']}" style="background-color:{C['page']};padding:18px 0 4px;border-top:1px solid {C['border']};">
    <div style="font-size:9px;color:{C['dim']};letter-spacing:0.3px;line-height:1.6;">
      QRadar Rule Status Auditor &nbsp;·&nbsp; Auto-generated {run_time}<br>
      Full rule details, and the Rule Testing Tracker (if configured), are in the attached Excel workbook.
    </div>
  </td></tr>

</table>

</td>
</tr>
</table>
</body></html>"""


def create_outlook_draft(excel_path, subject, html_body, images, high_importance=False):
    try:
        outlook = win32com.client.Dispatch('Outlook.Application')
        mail    = outlook.CreateItem(0)
        mail.Subject = subject
        if high_importance:
            mail.Importance = 2

        if excel_path and os.path.exists(excel_path):
            mail.Attachments.Add(excel_path)

        for cid, img_path in images.items():
            if img_path and os.path.exists(img_path):
                att = mail.Attachments.Add(img_path)
                att.PropertyAccessor.SetProperty(_MAPI_PR_ATTACH_CONTENT_ID, cid)

        mail.HTMLBody = html_body
        mail.Display()
        print("\n✉  Outlook draft created.")
    except Exception as e:
        logger.error("Outlook draft failed:\n%s", traceback.format_exc())
        print(f"\n❌ Outlook draft error: {e}")
    finally:
        for cid, img_path in images.items():
            if img_path and os.path.exists(img_path):
                try:
                    os.remove(img_path)
                except Exception:
                    pass


def main():
    if not VERIFY_SSL:
        urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

    print("=" * 62)
    print(f"  QRadar — {REPORT_TITLE}")
    print("=" * 62)
    print(f"  Host                     : {QRADAR_HOST}")
    print(f"  UCM app ID               : {UCM_APP_ID or '(not configured)'}")
    print(f"  UCM auth                 : {'SEC token' if UCM_AUTH_TOKEN else 'Basic Auth fallback'}")
    print(f"  Newly-triggered window   : last {NEWLY_TRIGGERED_WINDOW_DAYS} day(s)")
    print(f"  Tracker validation window: last {TRACKER_VALIDATION_WINDOW_DAYS} day(s)")
    print(f"  Tracker file             : {TRACKER_EXCEL_PATH or '(not configured)'}")
    print(f"  Tracker sheet            : {TRACKER_SHEET_NAME or '(first sheet)'}")
    print(f"  Auto-mark Triggered      : {AUTO_UPDATE_TRACKER_TRIGGERED}")
    print(f"  UCM_DEBUG                : {UCM_DEBUG}")
    print("=" * 62)

    if not validate_config():
        print("❌ Fix the configuration issues above before running.")
        return

    if not test_connection():
        return

    try:
        ucm_rows = fetch_ucm_rule_report()
    except Exception as e:
        print(f"\n❌ Use Case Manager fetch failed: {e}")
        print("   Set UCM_DEBUG = True and re-run to see the raw request/response and pinpoint why.")
        return

    if not ucm_rows:
        print("\n❌ Use Case Explorer returned no rows. Check UCM_APP_ID, UCM_AUTH_TOKEN, and that "
              "the filters in fetch_ucm_rule_report() match your deployment (UCM_DEBUG = True "
              "shows the raw request/response).")
        return

    print("\n🔍 Classifying rules...")
    master_df = build_master_rule_table(ucm_rows)
    dead_df   = master_df[master_df['is_dead']].copy()
    newly_df  = master_df[master_df['is_newly_triggered']].copy()

    print(f"   Total in-scope rules: {len(master_df)}")
    print(f"   Dead — never triggered: {len(dead_df)}")
    print(f"   Newly triggered (first fired in last {NEWLY_TRIGGERED_WINDOW_DAYS}d): {len(newly_df)}")

    print("\n📋 Loading rule testing tracker...")
    tracker_info = load_tracker_file(TRACKER_EXCEL_PATH) if TRACKER_EXCEL_PATH else None
    if tracker_info is None:
        print("   ℹ  Tracker not configured or unavailable this run.")
    else:
        tracker_info = reconcile_tracker_with_qradar(tracker_info, master_df)
        print(f"   ✅ {tracker_info['under_investigation_count']} under investigation, "
              f"{tracker_info['triggered_count']} triggered "
              f"(of {len(tracker_info['raw'])} tracked row(s), sheet '{tracker_info['sheet_name']}').")
        if tracker_info['name_col'] and tracker_info['unmatched_count']:
            print(f"   ⚠  {tracker_info['unmatched_count']} tracker row(s) didn't match any in-scope QRadar rule by name.")

        if tracker_info['changed_cells']:
            result = persist_tracker_updates(TRACKER_EXCEL_PATH, tracker_info)
            n_auto = len(tracker_info['auto_marked_rows'])
            if result == 'written':
                print(f"   ✏  Auto-marked {n_auto} rule(s) as Triggered — written back to {TRACKER_EXCEL_PATH} "
                      f"(backup saved alongside it).")
            elif result == 'locked':
                print(f"   ⚠  {n_auto} rule(s) qualify to be auto-marked Triggered, but the tracker file is "
                      f"open elsewhere — close it and re-run to persist.")
            elif result == 'disabled':
                print(f"   ℹ  {n_auto} rule(s) qualify to be auto-marked Triggered "
                      f"(AUTO_UPDATE_TRACKER_TRIGGERED=False — not written back).")
            elif result == 'error':
                print("   ❌ Could not write updates back to the tracker file — see log for details.")

    print(f"\n💾 Saving Excel report → {OUTPUT_EXCEL}")
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    saved_excel_path = save_excel_report(dead_df, newly_df, tracker_info, len(master_df), OUTPUT_EXCEL)

    print("\n📊 Generating status chart...")
    chart_path = generate_status_chart(len(dead_df), len(newly_df), tracker_info, len(master_df))

    print("\n✉  Building email draft...")
    html_body = build_email_html(dead_df, newly_df, tracker_info, len(master_df),
                                  chart_cid='rule_status' if chart_path else None)

    subject_parts = [f"{len(dead_df)} dead", f"{len(newly_df)} newly triggered"]
    if tracker_info is not None:
        subject_parts.append(f"{tracker_info['under_investigation_count']} under investigation")
        subject_parts.append(f"{tracker_info['triggered_count']} triggered")
    subject = f"{REPORT_TITLE} — " + ", ".join(subject_parts)

    high_importance = len(dead_df) >= EMAIL_HIGH_IMPORTANCE_DEAD_THRESHOLD

    images = {'rule_status': chart_path} if chart_path else {}
    create_outlook_draft(saved_excel_path, subject, html_body, images, high_importance=high_importance)
    print("\n✅ Done!")


if __name__ == '__main__':
    main()
