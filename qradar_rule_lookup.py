#!/usr/bin/env python3
"""
qradar_rule_lookup.py
---------------------
Reads Rule IDs from an Excel sheet, asks QRadar which offenses each rule
contributed to, and writes the results back into the same sheet.

Columns added/filled (created next to your data if they don't exist yet):
    Rule Name | Rule Enabled | Offense Count | Last Offense ID |
    Last Offense Description | Last Offense Status | Last Offense Magnitude |
    First Triggered | Last Triggered | Last Checked | Lookup Status

Usage:
    python qradar_rule_lookup.py                      # uses EXCEL_FILE below
    python qradar_rule_lookup.py my_rules.xlsx --days 180
    python qradar_rule_lookup.py my_rules.xlsx --dry-run   # print, don't save
"""
import argparse
import os
import re
import shutil
import sys
from copy import copy
from datetime import datetime
from pathlib import Path

import requests
import urllib3
from openpyxl import load_workbook
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

# ─── QRadar ───────────────────────────────────────────────────────────────────
QRADAR_HOST     = os.environ.get('QRADAR_HOST',     'https://your-qradar-host')
QRADAR_USERNAME = os.environ.get('QRADAR_USERNAME', 'your-username')
QRADAR_PASSWORD = os.environ.get('QRADAR_PASSWORD', 'your-password')
QRADAR_SEC_TOKEN = os.environ.get('QRADAR_SEC_TOKEN', '')   # optional; wins over user/pass
QRADAR_API_VERSION = os.environ.get('QRADAR_API_VERSION', '19.0')
VERIFY_SSL      = False          # set True + supply a CA bundle in production

# ─── Excel ────────────────────────────────────────────────────────────────────
EXCEL_FILE     = os.environ.get('EXCEL_FILE',     'rules.xlsx')
EXCEL_SHEET    = os.environ.get('EXCEL_SHEET',    '')          # blank = active sheet
RULE_ID_HEADER = os.environ.get('RULE_ID_HEADER', 'Rule ID')   # header text to look for
LOOKBACK_DAYS  = int(os.environ.get('LOOKBACK_DAYS', '90'))    # how far back to search
DISPLAY_TZ     = os.environ.get('DISPLAY_TZ',     'UTC')       # e.g. 'Asia/Kolkata'
PAGE_SIZE      = int(os.environ.get('PAGE_SIZE',  '200'))      # QRadar Range page size
MAKE_BACKUP    = os.environ.get('MAKE_BACKUP', 'true').lower() != 'false'

OUT_COLUMNS = [
    'Rule Name', 'Rule Enabled', 'Offense Count', 'Last Offense ID',
    'Last Offense Description', 'Last Offense Status', 'Last Offense Magnitude',
    'First Triggered', 'Last Triggered', 'Last Checked', 'Lookup Status',
]
DATE_COLUMNS = {'First Triggered', 'Last Triggered', 'Last Checked'}
DATE_FORMAT = 'yyyy-mm-dd hh:mm:ss'
UUID_RE = re.compile(r'^[0-9a-fA-F]{8}(-[0-9a-fA-F]{4}){3}-[0-9a-fA-F]{12}$')


# ─── Helpers ──────────────────────────────────────────────────────────────────
def die(msg: str):
    print(f'ERROR: {msg}', file=sys.stderr)
    sys.exit(1)


def check_config():
    if 'your-qradar-host' in QRADAR_HOST:
        die('QRADAR_HOST is still the placeholder. Set the QRADAR_HOST env var.')
    if not QRADAR_SEC_TOKEN and ('your-username' in QRADAR_USERNAME
                                 or 'your-password' in QRADAR_PASSWORD):
        die('Set QRADAR_SEC_TOKEN, or both QRADAR_USERNAME and QRADAR_PASSWORD.')
    if not VERIFY_SSL:
        urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)


def get_tz():
    try:
        from zoneinfo import ZoneInfo
        return ZoneInfo(DISPLAY_TZ)
    except Exception:
        print(f'WARN: timezone "{DISPLAY_TZ}" unavailable (pip install tzdata on '
              'Windows); falling back to UTC.', file=sys.stderr)
        from datetime import timezone
        return timezone.utc


def ms_to_dt(ms, tz):
    """Epoch milliseconds -> naive datetime in display tz (Excel can't store tz)."""
    return datetime.fromtimestamp(ms / 1000, tz=tz).replace(tzinfo=None)


def make_session() -> requests.Session:
    s = requests.Session()
    s.headers.update({'Accept': 'application/json', 'Version': QRADAR_API_VERSION})
    if QRADAR_SEC_TOKEN:
        s.headers['SEC'] = QRADAR_SEC_TOKEN
    else:
        s.auth = (QRADAR_USERNAME, QRADAR_PASSWORD)
    s.verify = VERIFY_SSL
    retry = Retry(total=4, backoff_factor=1.5, allowed_methods=('GET',),
                  status_forcelist=(429, 500, 502, 503, 504))
    s.mount('https://', HTTPAdapter(max_retries=retry))
    s.mount('http://', HTTPAdapter(max_retries=retry))
    return s


def get_all(session, path, params):
    """GET a QRadar collection, following Range-header pagination."""
    url = f"{QRADAR_HOST.rstrip('/')}/api{path}"
    results, start = [], 0
    while True:
        r = session.get(url, params=params, timeout=90,
                        headers={'Range': f'items={start}-{start + PAGE_SIZE - 1}'})
        if r.status_code == 401:
            die('QRadar returned 401 - check credentials / SEC token.')
        if r.status_code == 403:
            die(f'QRadar returned 403 for {path} - the account/token lacks permission.')
        r.raise_for_status()
        page = r.json()
        results.extend(page)
        start += len(page)
        m = re.search(r'/(\d+)$', r.headers.get('Content-Range', ''))
        total = int(m.group(1)) if m else None
        if not page or (total is not None and start >= total) \
                or (total is None and len(page) < PAGE_SIZE):
            return results


# ─── QRadar lookups ───────────────────────────────────────────────────────────
def load_rules(session):
    rules = get_all(session, '/analytics/rules',
                    {'fields': 'id,identifier,name,enabled,type'})
    by_id = {str(r['id']): r for r in rules}
    by_uuid = {str(r.get('identifier', '')).lower(): r for r in rules}
    print(f'Loaded {len(rules)} rules from QRadar.')
    return by_id, by_uuid


def resolve_rule(raw, by_id, by_uuid):
    """Sheet value -> QRadar rule dict, or None. Accepts numeric id or UUID."""
    if raw is None:
        return None
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    key = str(raw).strip()
    if UUID_RE.match(key):
        return by_uuid.get(key.lower())
    return by_id.get(key)


def load_offenses(session, days):
    cutoff_ms = int((datetime.now().timestamp() - days * 86400) * 1000)
    offenses = get_all(session, '/siem/offenses', {
        'fields': 'id,description,status,magnitude,start_time,last_updated_time,rules',
        'filter': f'last_updated_time > {cutoff_ms}',
    })
    print(f'Loaded {len(offenses)} offenses updated in the last {days} days.')
    return offenses


def aggregate(offenses, wanted_ids):
    """rule id -> {count, first_start, last_offense}"""
    stats = {}
    for o in offenses:
        for rid in {ru.get('id') for ru in (o.get('rules') or [])}:
            if rid not in wanted_ids:
                continue
            s = stats.get(rid)
            if s is None:
                stats[rid] = {'count': 1, 'first_start': o['start_time'], 'last': o}
                continue
            s['count'] += 1
            s['first_start'] = min(s['first_start'], o['start_time'])
            if o['last_updated_time'] > s['last']['last_updated_time']:
                s['last'] = o
    return stats


# ─── Excel ────────────────────────────────────────────────────────────────────
def find_header(ws):
    """Locate the Rule ID header within the first 15 rows."""
    target = RULE_ID_HEADER.strip().lower()
    for row in ws.iter_rows(min_row=1, max_row=15):
        for c in row:
            if c.value is not None and str(c.value).strip().lower() == target:
                return c.row, c.column
    die(f'Could not find a "{RULE_ID_HEADER}" header in the first 15 rows of '
        f'sheet "{ws.title}".')


def map_output_columns(ws, header_row, id_col):
    """Reuse existing output headers; append missing ones after the last header."""
    existing = {}
    last_col = 0
    for c in ws[header_row]:
        if c.value is not None and str(c.value).strip():
            existing[str(c.value).strip().lower()] = c.column
            last_col = max(last_col, c.column)
    style_src = ws.cell(header_row, id_col)
    cols = {}
    for name in OUT_COLUMNS:
        if name.lower() in existing:
            cols[name] = existing[name.lower()]
            continue
        last_col += 1
        cell = ws.cell(header_row, last_col, name)
        cell.font, cell.fill = copy(style_src.font), copy(style_src.fill)
        cell.border, cell.alignment = copy(style_src.border), copy(style_src.alignment)
        ws.column_dimensions[cell.column_letter].width = \
            50 if name == 'Last Offense Description' else 22
        cols[name] = last_col
    return cols


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('file', nargs='?', default=EXCEL_FILE)
    ap.add_argument('--days', type=int, default=LOOKBACK_DAYS)
    ap.add_argument('--dry-run', action='store_true', help='print results, do not save')
    args = ap.parse_args()

    check_config()
    path = Path(args.file)
    if not path.exists():
        die(f'Excel file not found: {path}')

    wb = load_workbook(path, keep_vba=path.suffix.lower() == '.xlsm')
    ws = wb[EXCEL_SHEET] if EXCEL_SHEET else wb.active
    header_row, id_col = find_header(ws)
    cols = map_output_columns(ws, header_row, id_col)

    # Collect rows that have a Rule ID
    rows = [(r, ws.cell(r, id_col).value) for r in range(header_row + 1, ws.max_row + 1)
            if ws.cell(r, id_col).value not in (None, '')]
    if not rows:
        die('No Rule IDs found under the header.')
    print(f'Found {len(rows)} rule rows in sheet "{ws.title}".')

    session = make_session()
    by_id, by_uuid = load_rules(session)
    resolved = {r: resolve_rule(v, by_id, by_uuid) for r, v in rows}
    wanted = {rule['id'] for rule in resolved.values() if rule}
    stats = aggregate(load_offenses(session, args.days), wanted)

    tz = get_tz()
    now = datetime.now(tz).replace(tzinfo=None)
    for r, raw in rows:
        rule = resolved[r]
        out = {n: None for n in OUT_COLUMNS}
        out['Last Checked'] = now
        if rule is None:
            out['Lookup Status'] = 'Rule ID not found in QRadar'
        else:
            out['Rule Name'] = rule.get('name')
            out['Rule Enabled'] = 'Yes' if rule.get('enabled') else 'No'
            s = stats.get(rule['id'])
            if s is None:
                out['Offense Count'] = 0
                out['Lookup Status'] = f'No offenses in last {args.days} days'
            else:
                o = s['last']
                out.update({
                    'Offense Count': s['count'],
                    'Last Offense ID': o['id'],
                    'Last Offense Description': (o.get('description') or '').strip(),
                    'Last Offense Status': o.get('status'),
                    'Last Offense Magnitude': o.get('magnitude'),
                    'First Triggered': ms_to_dt(s['first_start'], tz),
                    'Last Triggered': ms_to_dt(o['last_updated_time'], tz),
                    'Lookup Status': f'OK (last {args.days} days)',
                })
        print(f"  row {r:>4}  rule {str(raw):<12} -> {out['Lookup Status']}"
              + (f"  offense #{out['Last Offense ID']}" if out['Last Offense ID'] else ''))
        for name, value in out.items():
            cell = ws.cell(r, cols[name], value)
            if name in DATE_COLUMNS and value is not None:
                cell.number_format = DATE_FORMAT

    if args.dry_run:
        print('Dry run - nothing saved.')
        return
    if MAKE_BACKUP:
        backup = path.with_name(f'{path.stem}.bak{path.suffix}')
        shutil.copy2(path, backup)
        print(f'Backup written: {backup}')
    try:
        wb.save(path)
    except PermissionError:
        die(f'Cannot write {path}. Close it in Excel and run again.')
    print(f'Updated {path}')


if __name__ == '__main__':
    main()
